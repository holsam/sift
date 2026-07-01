'''
Sift UI: preview panel
'''

# -- Import external dependencies --
from pathlib import Path
from PySide6.QtCore import Qt, QUrl
from PySide6.QtGui import QMovie, QPixmap
from PySide6.QtMultimedia import QAudioOutput, QMediaPlayer
from PySide6.QtMultimediaWidgets import QVideoWidget
from PySide6.QtPdf import QPdfDocument
from PySide6.QtPdfWidgets import QPdfView
from PySide6.QtWidgets import (
    QLabel,
    QPlainTextEdit,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)

# -- Import internal classes --
from sift.ui.utils.media_controls import MediaControls

# -- Define filetype extensions --
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".tiff"}
VIDEO_EXTS = {".mp4", ".mov", ".mkv", ".avi", ".webm"}
AUDIO_EXTS = {".mp3", ".wav", ".flac", ".ogg", ".m4a", ".aac"}
TEXT_EXTS = {
    ".txt", ".md", ".py", ".js", ".ts", ".json", ".csv", ".tsv", ".log",
    ".html", ".css", ".yaml", ".yml", ".xml", ".ini", ".cfg", ".sh", ".gd",
}

# -- Define character limit for reading text --
TEXT_READ_LIMIT = 200_000

# -- PreviewWidget: class to define structure of preview panel and logic for swapping content to match filetype
class PreviewWidget(QWidget):
    # Initialise QWidget
    def __init__(self) -> None:
        super().__init__()
        self._layout = QVBoxLayout(self)
        self._player: QMediaPlayer | None = None
        self._audio: QAudioOutput | None = None
        self.show_message('Nothing to preview')

    # show_file: shows a file in the preview panel
    def show_file(self, path: Path) -> None:
        self._teardown_media()
        self._clear()
        ext = path.suffix.lower()
        if ext in IMAGE_EXTS:
            self._show_image(path)
        elif ext in VIDEO_EXTS:
            self._show_av(path, video=True)
        elif ext in AUDIO_EXTS:
            self._show_av(path, video=False)
        elif ext == '.pdf':
            self._show_pdf(path)
        elif ext == '.docx':
            self._show_text(_docx_text(path))
        elif ext == '.xlsx':
            self._show_table(_xlsx_rows(path))
        elif ext == '.pptx':
            self._show_text(_pptx_text(path))
        elif ext in TEXT_EXTS:
            self._show_text(_read_text(path))
        else:
            self.show_message(f'No preview available for {ext or 'this file type'}')

    # show_message: shows a message in the preview panel
    def show_message(self, text: str) -> None:
        self._teardown_media()
        self._clear()
        label = QLabel(text, alignment=Qt.AlignmentFlag.AlignCenter)
        self._layout.addWidget(label)

    # _show_image: displays an image (or movie if GIF)
    def _show_image(self, path: Path) -> None:
        label = QLabel(alignment=Qt.AlignmentFlag.AlignCenter)
        if path.suffix.lower() == ".gif":
            movie = QMovie(str(path))
            label.setMovie(movie)
            movie.start()
        else:
            pixmap = QPixmap(str(path))
            label.setPixmap(
                pixmap.scaled(
                    self.size(),
                    Qt.AspectRatioMode.KeepAspectRatio,
                    Qt.TransformationMode.SmoothTransformation,
                )
            )
        self._layout.addWidget(label)

    # _show_av: displays a video or audio file
    def _show_av(self, path: Path, video: bool) -> None:
        self._player = QMediaPlayer()
        self._audio = QAudioOutput()
        self._player.setAudioOutput(self._audio)
        if video:
            surface = QVideoWidget()
            self._player.setVideoOutput(surface)
            self._layout.addWidget(surface, stretch=1)
        else:
            self._layout.addWidget(
                QLabel(path.name, alignment=Qt.AlignmentFlag.AlignCenter),
                stretch=1,
            )
        self._player.setSource(QUrl.fromLocalFile(str(path)))
        self._layout.addWidget(MediaControls(self._player, self._audio))

    # _show_pdf: displays a PDF file
    def _show_pdf(self, path: Path) -> None:
        doc = QPdfDocument(self)
        doc.load(str(path))
        view = QPdfView()
        view.setDocument(doc)
        view.setPageMode(QPdfView.PageMode.MultiPage)
        self._layout.addWidget(view)

    # _show_text: displays a text file
    def _show_text(self, text: str) -> None:
        editor = QPlainTextEdit()
        editor.setReadOnly(True)
        editor.setPlainText(text)
        self._layout.addWidget(editor)

    # _show_table: displays a spreadsheet as a table
    def _show_table(self, rows: list[list[str]]) -> None:
        if not rows:
            self.show_message('Empty spreadsheet')
            return
        cols = max(len(r) for r in rows)
        table = QTableWidget(len(rows), cols)
        for r, row in enumerate(rows):
            for c in range(cols):
                value = row[c] if c < len(row) else ''
                table.setItem(r, c, QTableWidgetItem(value))
        self._layout.addWidget(table)

    # _clear: clears the current preview
    def _clear(self) -> None:
        while self._layout.count():
            item = self._layout.takeAt(0)
            w = item.widget()
            if w:
                w.deleteLater()

    # _teardown_media: stops any current media previews
    def _teardown_media(self) -> None:
        if self._player is not None:
            self._player.stop()
            self._player.setVideoOutput(None)
            self._player = None
            self._audio = None

# -- _read_text: reads a text file to a string -- 
def _read_text(path: Path) -> str:
    try:
        with path.open('r', encoding='utf-8', errors='replace') as f:
            return f.read(TEXT_READ_LIMIT)
    except OSError as exc:
        return f'Could not read file: {exc}'

# -- _docx_text: read an Office Word document to a string --
def _docx_text(path: Path) -> str:
    try:
        import docx
    except ImportError:
        return 'python-docx not installed'
    document = docx.Document(str(path))
    return '\n'.join(p.text for p in document.paragraphs)

# -- _xlsx_rows: read an Office Excel spreadsheet to a list of lists of strings --
def _xlsx_rows(path: Path) -> list[list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return [['openpyxl not installed']]
    wb = load_workbook(str(path), read_only=True, data_only=True)
    ws = wb.active
    rows: list[list[str]] = []
    for row in ws.iter_rows(values_only=True):
        rows.append(['' if v is None else str(v) for v in row])
        if len(rows) >= 500:  # cap for responsiveness
            break
    wb.close()
    return rows

# -- _pptx_text: read an Office Powerpoint presentation to a string --
def _pptx_text(path: Path) -> str:
    try:
        from pptx import Presentation
    except ImportError:
        return 'python-pptx not installed'
    prs = Presentation(str(path))
    chunks: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        chunks.append(f'--- Slide {i} ---')
        for shape in slide.shapes:
            if shape.has_text_frame:
                chunks.append(shape.text_frame.text)
    return '\n'.join(chunks)